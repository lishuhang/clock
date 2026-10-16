#!/usr/bin/env python3
"""v15 fix A: 修复v15-stageA.docx的六处缺陷 → v15-stageA2.docx
1.标题6 17→18地区(最终裁决:10-14线上手册18>10-05手册17,v14已定18)
2.段103喵影句:修复重复子句+找回被误删的LibTV/TapNow智能剪辑句(v14黄段)
3.表S电商列连锁: Seko 8→5、AniShort 5→6、纳逗Pro 6→7、TapNow 7→8(修正#1,StageA漏改)
4.表4(电商需求点)对齐最终草稿: LibTV协作格×→√(StageA漏改)+小云雀18地区+措辞统一
5.剪辑表万兴行: 按批注8原文"+导入喵影"
6.独家表: 小云雀18地区 + 纳逗专业影视格去重(A5 cell_lines未清旧段落bug)
"""
import copy
import docx
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = '/home/z/my-project/materials/下篇-1015/v15-stageA.docx'
OUT = '/home/z/my-project/materials/下篇-1015/v15-stageA2.docx'
d = docx.Document(SRC)
LOG = []

def get_rpr_tmpl(el):
    for r in el.iter(qn('w:r')):
        if r.find(qn('w:t')) is not None and (r.find(qn('w:t')).text or '').strip():
            rpr = r.find(qn('w:rPr'))
            return copy.deepcopy(rpr) if rpr is not None else None
    return None

def make_run(text, rpr_tmpl, highlight=False):
    r = OxmlElement('w:r')
    if rpr_tmpl is not None:
        r.append(copy.deepcopy(rpr_tmpl))
    if highlight:
        rpr = r.find(qn('w:rPr'))
        if rpr is None:
            rpr = OxmlElement('w:rPr'); r.insert(0, rpr)
        hl = OxmlElement('w:highlight'); hl.set(qn('w:val'), 'yellow')
        rpr.append(hl)
    t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = text
    r.append(t)
    return r

def shade_cell(cell, fill='FFFF00'):
    tcPr = cell._tc.get_or_add_tcPr()
    for shd in tcPr.findall(qn('w:shd')):
        tcPr.remove(shd)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), fill)
    tcPr.append(shd)

def set_cell_text(cell, text, shade=False):
    rpr_tmpl = get_rpr_tmpl(cell._tc)
    first = cell.paragraphs[0]
    for child in list(first._p):
        if child.tag != qn('w:pPr'):
            first._p.remove(child)
    first._p.append(make_run(text, rpr_tmpl))
    for p in cell.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    if shade:
        shade_cell(cell)

def cell_lines(cell, lines, shade=False):
    """清空全部旧段落后写入多行(修复StageA版未清旧段落的bug)"""
    rpr_tmpl = get_rpr_tmpl(cell._tc)
    first = cell.paragraphs[0]
    pPr0 = first._p.find(qn('w:pPr'))
    for p in cell.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    for child in list(first._p):
        if child.tag != qn('w:pPr'):
            first._p.remove(child)
    first._p.append(make_run(lines[0], rpr_tmpl))
    prev = first._p
    for ln in lines[1:]:
        newp = OxmlElement('w:p')
        if pPr0 is not None:
            newp.append(copy.deepcopy(pPr0))
        newp.append(make_run(ln, rpr_tmpl))
        prev.addnext(newp); prev = newp
    if shade:
        shade_cell(cell)

def replace_para_runs(para, segments):
    rpr_tmpl = get_rpr_tmpl(para._p)
    for child in list(para._p):
        if child.tag != qn('w:pPr'):
            para._p.remove(child)
    for txt, hl in segments:
        para._p.append(make_run(txt, rpr_tmpl, highlight=hl))

def find_para(substr):
    for p in d.paragraphs:
        if substr in p.text:
            return p
    return None

# ═══ F1. 标题6: 17→18 ═══
p6 = find_para('6. 标题：一键出海17地区')
if p6 is not None:
    replace_para_runs(p6, [('6. 标题：一键出海18地区、6语种转绘，出海谁最强丨AI视频画布横评·功能篇　摘要：整部剧换国家重拍，转绘流水线卷起来', False)])
    LOG.append('F1 标题6: 17地区→18地区(最终裁决)')

# ═══ F2. 段103 喵影句重建(修重复+找回LibTV/TapNow句) ═══
p_wx = find_para('抽卡成片之后')
if p_wx is not None:
    replace_para_runs(p_wx, [
        ('抽卡成片之后，各家平台的下一步操作不尽相同，有的可以导出到其他专业编辑器，比如小云雀导出到剪映；', False),
        ('有的联动自家剪辑工具，比如万兴剧厂联动喵影——画布内只能裁剪长短（实测），但定稿视频与字幕可以一键同步到万兴喵影客户端继续剪辑，缺素材无需本地导出，后期修改双端实时同步', True),
        ('；有的在浏览器做了编辑器，比如Seko、AniShort；', False),
        ('还有的把剪辑自动化——LibTV的智能剪辑按口播、批量广告、混剪等场景，输入文稿或一句话主题自动成片；TapNow的智能剪辑则自动识别镜头切换、把成片拆成分镜', True),
        ('。', False),
    ])
    LOG.append('F2 喵影句: 重建(修复StageA重复子句bug,找回LibTV/TapNow智能剪辑v14黄段)')

# ═══ F3. 表S 电商列连锁 ═══
tS = d.tables[0]
NEW_ECOM = {5: 'Seko', 6: 'AniShort', 7: '纳逗Pro', 8: 'TapNow'}
for rank, plat in NEW_ECOM.items():
    cell = tS.rows[rank].cells[3]
    if cell.text.strip() != plat:
        set_cell_text(cell, plat, shade=True)
        LOG.append(f'F3 表S电商列 第{rank}名: {cell.text.strip()}→{plat}(黄)')
    else:
        set_cell_text(cell, plat, shade=False)

# ═══ F4. 表4(电商需求点) 对齐最终草稿 ═══
t4 = d.tables[3]
UPD = {
    'OiiOii':   {4: '√ 20子账号+子账号/小组积分限额+团队版席位管理'},
    '小云雀':   {2: '√ 重制转绘18地区（规模最大）',
                 4: '√ 团队版4档：画布实时协作+10-20席位+分组隔离+角色权限'},
    'LibTV':    {4: '√ 高级/超级团队会员：多人画布协作+故事板协作+席位管理+积分灵活分配'},
    '万兴剧厂': {4: '√ 团队版三档（入门/专业/企业）：多账号创建角色+权限分级+增购席位+用量管控',
                 5: '√ MCP&CLI（CLI公测版）'},
    'AniShort': {5: '√ AniShort CLI（正式上线）+3D导演台MCP'},
    'Seko':     {4: '√ 团队标准/高级/至尊会员（4/8/15席位）+企业会员'},
}
for r in t4.rows[1:]:
    plat = r.cells[0].text.strip()
    if plat in UPD:
        for ci, txt in UPD[plat].items():
            old = r.cells[ci].text
            if old.strip() != txt:
                set_cell_text(r.cells[ci], txt, shade=True)
                LOG.append(f'F4 表4 {plat}列{ci}: {old[:20]}…→ {txt[:30]}…(黄)')

# ═══ F5. 剪辑表(表6) 万兴行 按批注8原文 ═══
tbian = d.tables[5]
for r in tbian.rows:
    if r.cells[0].text.strip() == '万兴剧厂':
        set_cell_text(r.cells[1], '画布内剪辑只能裁长短（实测）+导入喵影', shade=True)
        LOG.append('F5 剪辑表万兴行: →画布内剪辑只能裁长短（实测）+导入喵影(批注8原文,黄)')
        break

# ═══ F6. 独家表(表7): 小云雀18地区 + 纳逗去重 ═══
t7 = d.tables[6]
for r in t7.rows[1:]:
    plat = r.cells[0].text.strip()
    if plat == '小云雀':
        if '17地区' in r.cells[3].text:
            set_cell_text(r.cells[3], r.cells[3].text.replace('17地区', '18地区'), shade=True)
            LOG.append('F6 独家表小云雀电商格: 17→18地区(黄)')
    if plat == '纳逗Pro':
        txt = r.cells[2].text
        if '归发布者' in txt or txt.count('Skills定价上架') > 1:
            cell_lines(r.cells[2], ['• 6000+IP库+数字资产库（364角色图/168模型图/309视频/14个3D模型）',
                                    '• Skills定价上架、积分分成全部给到发布者'], shade=True)
            LOG.append('F6 独家表纳逗影视格: 去重(修复StageA cell_lines未清旧段落bug,黄)')

d.save(OUT)
print('\n'.join(LOG))
print('FIX A OK ->', OUT)
