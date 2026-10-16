#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""v15信息图: 从v15成品docx提取11张表 → skill-a(style-a)静态信息图 HTML+PNG
规范(用户三条): 1)字号默认28px不加大 2)表格结构严格依从原文(标题/表头/图例/脚注),内容按裁决修正 3)不标黄
"""
import copy as _copy
import glob, html as H, os, shutil, subprocess, sys
from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

DOCX = sorted(glob.glob('/home/z/my-project/materials/下篇-1015/漫剧智能体功能横评_功能篇_v15-*.docx'))[-1]
TMPL = Path('/home/z/my-project/materials/yz-infographic/yz-infographic-v3-1/style-a.html').read_text(encoding='utf-8')
ASSETS_SRC = '/home/z/my-project/materials/yz-infographic/yz-infographic-v3-1/assets'
OUT = Path('/home/z/my-project/download/v15交付/信息图-v15')
OUT.mkdir(parents=True, exist_ok=True)
if not (OUT / 'assets').exists():
    shutil.copytree(ASSETS_SRC, OUT / 'assets')

d = docx.Document(DOCX)

def iter_blocks(doc):
    for c in doc.element.body.iterchildren():
        if c.tag == qn('w:p'): yield Paragraph(c, doc)
        elif c.tag == qn('w:tbl'): yield Table(c, doc)
tables = [b for b in iter_blocks(d) if isinstance(b, Table)]

# ═══ 表元数据: (文件名, 标题[docx表标题], 图例/注[原文中表上方行], 列宽) ═══
META = [
    ('imgS',  '按用户画像分的各平台功能适配情况', None, ['10%', '30%', '30%', '30%']),
    ('img1', '漫剧连载类需求点', '√有　△有一部分/间接实现　×没有', ['19%', '17%', '16%', '16%', '16%', '16%']),
    ('img2', '专业影视类需求点', '√有　△有一部分/间接实现　×没有', ['18%', '17%', '17%', '16%', '16%', '16%']),
    ('img3', '电商出海类需求点', '√有　△有一部分/间接实现　×没有', ['17%', '18%', '16%', '15%', '17%', '17%']),
    ('img4', 'AI视频画布的"标配"', None, ['30%', '70%']),
    ('img5', '抽完后的剪辑动作', None, ['24%', '76%']),
    ('img6', '"人无我有"和独家功能', '注：有的功能可能面向多种人群，会被归类到它可能帮助最大的一群。', ['15%', '28%', '29%', '28%']),
    ('img7', '角色一致性及参考主体相关能力', None, ['22%', '78%']),
    ('img8', '剧本与分镜能力', None, ['22%', '78%']),
    ('img9', '团队协作能力', None, ['14%', '17%', '18%', '17%', '17%', '17%']),
    ('img10', '短剧出海转绘能力', None, ['13%', '20%', '19%', '24%', '24%']),
]
FOOT = '数据来源：娱乐资本论根据官方手册和实测整理，信息截至2026年10月5日，如与实际有差异，以实际呈现为准。'

CSS = """<style>
.bz-sub{font-size:24px;color:var(--s);margin:-24px 0 18px;line-height:1.35}
.bz-table{table-layout:fixed}
.bz-table th{font-size:26px;padding:10px 8px;color:var(--s);border-bottom:3px solid var(--l);line-height:1.25;white-space:normal;vertical-align:bottom}
.bz-table td{padding:10px 8px;line-height:1.3;word-break:break-word;vertical-align:top;font-size:28px}
.bz-p{font-weight:700}
.bz-yes{color:var(--ad);font-weight:700}
.bz-part{color:var(--t);font-weight:700}
.bz-no{color:#c9c4c2;font-weight:400}
.bz-dash{color:#c9c4c2}
</style>"""

def esc(s):
    return H.escape(s, quote=False)

def line_html(line):
    line = line.strip()
    if not line:
        return ''
    if line in ('-', '–'):
        return '<span class="bz-dash">-</span>'
    if line[0] in '√△×':
        cls = {'√': 'bz-yes', '△': 'bz-part', '×': 'bz-no'}[line[0]]
        rest = esc(line[1:]).lstrip()
        return f'<span class="{cls}">{line[0]}</span>{rest and " " + rest}'
    return esc(line)

def cell_html(cell):
    lines = [p.text for p in cell.paragraphs if p.text.strip()]
    return '<br>'.join(line_html(l) for l in lines) or '<span class="bz-dash">-</span>'

def fig_html(title, legend, tbl, colw):
    hdr = ''.join(f'<th>{esc(c.text.strip())}</th>' for c in tbl.rows[0].cells)
    trs = []
    for r in tbl.rows[1:]:
        tds = []
        for j, c in enumerate(r.cells):
            cls = ' class="bz-p"' if j == 0 else ''
            tds.append(f'<td{cls}>{cell_html(c)}</td>')
        trs.append(f'<tr>{"".join(tds)}</tr>')
    cg = '<colgroup>' + ''.join(f'<col style="width:{x}">' for x in colw) + '</colgroup>'
    sub = f'<div class="bz-sub">{esc(legend)}</div>' if legend else ''
    body = f'{CSS}{sub}<table class="bz-table">{cg}<thead><tr>{hdr}</tr></thead><tbody>{"".join(trs)}</tbody></table>'
    page = TMPL.replace('【TITLE_HERE】', esc(title))
    page = page.replace('<div class="yz-placeholder">【CHART_BODY_HERE】</div>', body)
    page = page.replace('<div class="chart-source">【SOURCE_HERE】</div>', f'<div class="chart-source">{esc(FOOT)}</div>')
    return page

# ═══ 生成11个HTML + source-data.csv ═══
csv_rows = ['fig,row,col,value']
for (name, title, legend, colw), tbl in zip(META, tables):
    assert len(tbl.columns) == len(colw), f'{name} 列数{len(tbl.columns)}≠{len(colw)}'
    page = fig_html(title, legend, tbl, colw)
    (OUT / f'{name}-styled.html').write_text(page, encoding='utf-8')
    hdrs = [c.text.strip() for c in tbl.rows[0].cells]
    for r in tbl.rows[1:]:
        row_label = r.cells[0].text.strip().replace(',', '、')
        for j, c in enumerate(r.cells[1:], start=1):
            val = ' / '.join(p.text.strip() for p in c.paragraphs if p.text.strip()).replace(',', '，')
            csv_rows.append(f'{name},{row_label},{hdrs[j]},{val}')
    print(f'{name}: {len(tbl.rows)-1}行×{len(tbl.columns)}列 已生成')

(OUT / 'source-data.csv').write_text('\n'.join(csv_rows), encoding='utf-8')
print(f'source-data.csv: {len(csv_rows)-1}行数据')
print('HTML生成完毕 ->', OUT)
